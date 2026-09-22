from flask import Flask, render_template, request, jsonify, session, redirect, url_for
from database import init_db, get_db, hp, gen_password
import csv, io, base64
from datetime import datetime, timedelta
import openpyxl, io as _io

app = Flask(__name__, static_folder="static", static_url_path="/static")
app.secret_key = "evacuai-secret-2024"
app.config['SESSION_PERMANENT'] = True
app.config['PERMANENT_SESSION_LIFETIME'] = timedelta(hours=24)

init_db()

def uid(): return session.get("user")
def now(): return datetime.now().strftime("%Y-%m-%d %H:%M:%S")

# ── Pages ─────────────────────────────────────────────────────────────────────
@app.route("/")
def index():
    if uid():
        return redirect(url_for("admin_page") if session.get("role")=="admin" else url_for("student_page"))
    return render_template("index.html")

@app.route("/student")
def student_page():
    if not uid() or session.get("role") not in ("student","teacher","staff"):
        return redirect(url_for("index"))
    return render_template("student.html", name=session.get("name"))

@app.route("/admin")
def admin_page():
    if not uid() or session.get("role")!="admin": return redirect(url_for("index"))
    return render_template("admin.html", name=session.get("name"))

@app.route("/admin/students")
def students_page():
    if not uid() or session.get("role")!="admin": return redirect(url_for("index"))
    return render_template("students.html", name=session.get("name"))

@app.route("/profile")
def profile_page():
    if not uid(): return redirect(url_for("index"))
    return render_template("profile.html", name=session.get("name"), role=session.get("role"), sid=session.get("sid"))

@app.route("/map")
def map_page():
    if not uid(): return redirect(url_for("index"))
    return render_template("map.html", name=session.get("name"), role=session.get("role"))

@app.route("/logout")
def logout_page():
    session.clear(); session.modified=True
    return redirect(url_for("index"))

# ── Auth ──────────────────────────────────────────────────────────────────────
@app.route("/api/login", methods=["POST"])
def login():
    data = request.get_json()
    sid  = data.get("student_id","").strip().upper()
    pw   = data.get("password","")
    if not sid or not pw: return jsonify({"error":"Please enter Student ID and password."}),400
    db   = get_db()
    user = db.execute("SELECT * FROM users WHERE student_id=?", (sid,)).fetchone()
    if not user:
        return jsonify({"error":"Student ID not found. Please contact your department registrar."}),404
    if user["password"] != hp(pw):
        return jsonify({"error":"Incorrect password."}),401
    session.update({"user":user["id"],"name":user["name"],"role":user["role"],"sid":sid})
    return jsonify({"role":user["role"],"name":user["name"],"must_change_pw":bool(user["must_change_pw"])})

@app.route("/api/logout", methods=["POST"])
def logout():
    session.clear(); session.modified=True
    return jsonify({"message":"Logged out."})

@app.route("/api/me")
def me():
    if not uid(): return jsonify({"error":"Not logged in"}),401
    return jsonify({"id":uid(),"name":session.get("name"),"role":session.get("role")})

@app.route("/api/change-password", methods=["POST"])
def change_password():
    if not uid(): return jsonify({"error":"Unauthorized"}),401
    data   = request.get_json()
    old_pw = data.get("old_password","")
    new_pw = data.get("new_password","")
    if not new_pw or len(new_pw)<6:
        return jsonify({"error":"New password must be at least 6 characters."}),400
    db   = get_db()
    user = db.execute("SELECT * FROM users WHERE id=?", (uid(),)).fetchone()
    if user["password"] != hp(old_pw):
        return jsonify({"error":"Current password is incorrect."}),401
    db.execute("UPDATE users SET password=?, must_change_pw=0 WHERE id=?", (hp(new_pw), uid()))
    db.commit()
    return jsonify({"message":"Password changed successfully."})

# ── Users ─────────────────────────────────────────────────────────────────────
@app.route("/api/users", methods=["GET"])
def get_users():
    if not uid() or session.get("role")!="admin": return jsonify({"error":"Unauthorized"}),401
    db   = get_db()
    rows = db.execute("SELECT id,name,email,student_id,role,must_change_pw FROM users ORDER BY role,name").fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/users", methods=["POST"])
def add_user():
    if not uid() or session.get("role")!="admin": return jsonify({"error":"Unauthorized"}),401
    data = request.get_json()
    name = data.get("name","").strip()
    sid  = data.get("student_id","").strip()
    role = data.get("role","student")
    if not name or not sid: return jsonify({"error":"Name and Student ID are required."}),400
    if not sid.isdigit() or len(sid)!=8:
        return jsonify({"error":"Student ID must be exactly 8 digits."}),400
    pw = data.get("password","").strip() or gen_password(sid)
    db = get_db()
    if db.execute("SELECT id FROM users WHERE student_id=?", (sid,)).fetchone():
        return jsonify({"error":"Student ID already exists."}),409
    db.execute("INSERT INTO users (name,student_id,password,role,must_change_pw) VALUES (?,?,?,?,1)",
               (name, sid, hp(pw), role))
    db.commit()
    return jsonify({"message":"Student account created.", "password":pw})

@app.route("/api/users/upload-csv", methods=["POST"])
def upload_csv():
    if not uid() or session.get("role")!="admin": return jsonify({"error":"Unauthorized"}),401
    data     = request.get_json()
    csv_text = data.get("csv","")
    reader   = csv.DictReader(io.StringIO(csv_text))
    db       = get_db()
    added, skipped = 0, 0
    for row in reader:
        sid  = row.get("student_id","").strip()
        name = row.get("name","").strip()
        role = row.get("role","student").strip() or "student"
        pw   = row.get("password","").strip() or gen_password(sid)
        if not sid or not name: skipped+=1; continue
        if db.execute("SELECT id FROM users WHERE student_id=?", (sid,)).fetchone():
            skipped+=1; continue
        db.execute("INSERT INTO users (name,student_id,password,role,must_change_pw) VALUES (?,?,?,?,1)",
                   (name, sid, hp(pw), role))
        added+=1
    db.commit()
    return jsonify({"message":f"Added {added} accounts. Skipped {skipped} (duplicate or incomplete)."})

@app.route("/api/users/<int:user_id>", methods=["PATCH"])
def update_user(user_id):
    if not uid() or session.get("role")!="admin": return jsonify({"error":"Unauthorized"}),401
    data = request.get_json()
    db   = get_db()
    user = db.execute("SELECT * FROM users WHERE id=?", (user_id,)).fetchone()
    if not user: return jsonify({"error":"User not found."}),404
    name   = data.get("name", user["name"]).strip()
    role   = data.get("role", user["role"])
    new_pw = data.get("password","").strip()
    if new_pw:
        db.execute("UPDATE users SET name=?,role=?,password=?,must_change_pw=1 WHERE id=?",
                   (name, role, hp(new_pw), user_id))
    else:
        db.execute("UPDATE users SET name=?,role=? WHERE id=?", (name, role, user_id))
    db.commit()
    return jsonify({"message":"User updated."})

@app.route("/api/users/<int:user_id>", methods=["DELETE"])
def delete_user(user_id):
    if not uid() or session.get("role")!="admin": return jsonify({"error":"Unauthorized"}),401
    if user_id==uid(): return jsonify({"error":"Cannot delete your own account."}),400
    db = get_db()
    db.execute("DELETE FROM users WHERE id=?", (user_id,))
    db.execute("DELETE FROM safety_status WHERE user_id=?", (user_id,))
    db.execute("DELETE FROM reports WHERE user_id=?", (user_id,))
    db.commit()
    return jsonify({"message":"User deleted."})

# ── Alerts ────────────────────────────────────────────────────────────────────
@app.route("/api/alerts", methods=["GET"])
def get_alerts():
    db   = get_db()
    rows = db.execute("SELECT * FROM alerts WHERE status='approved' ORDER BY created_at DESC LIMIT 30").fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/alerts/pending", methods=["GET"])
def get_pending():
    if not uid() or session.get("role")!="admin": return jsonify({"error":"Unauthorized"}),401
    db   = get_db()
    rows = db.execute("SELECT * FROM alerts WHERE status='pending' ORDER BY created_at ASC").fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/alerts", methods=["POST"])
def post_alert():
    if not uid(): return jsonify({"error":"Unauthorized"}),401
    data    = request.get_json()
    kind    = data.get("type","general")
    message = data.get("message","").strip()
    role    = session.get("role","student")
    if not message: return jsonify({"error":"Message cannot be empty."}),400
    status  = "approved" if role=="admin" else "pending"
    db      = get_db()
    cur     = db.execute(
        "INSERT INTO alerts (type,message,sender,sender_role,created_at,user_id,status) VALUES (?,?,?,?,?,?,?)",
        (kind, message, session.get("name","Unknown"), role, now(), uid(), status)
    )
    db.commit()
    label = "Alert broadcast!" if role=="admin" else "Report submitted — waiting for admin approval."
    return jsonify({"message":label,"id":cur.lastrowid,"status":status})

@app.route("/api/alerts/<int:aid>", methods=["PATCH"])
def update_alert(aid):
    if not uid() or session.get("role")!="admin": return jsonify({"error":"Unauthorized"}),401
    data  = request.get_json()
    db    = get_db()
    alert = db.execute("SELECT * FROM alerts WHERE id=?", (aid,)).fetchone()
    if not alert: return jsonify({"error":"Alert not found."}),404
    msg  = data.get("message", alert["message"]).strip()
    kind = data.get("type", alert["type"])
    db.execute("UPDATE alerts SET message=?,type=? WHERE id=?", (msg, kind, aid))
    db.commit()
    return jsonify({"message":"Alert updated."})

@app.route("/api/alerts/<int:aid>", methods=["DELETE"])
def delete_alert(aid):
    if not uid(): return jsonify({"error":"Unauthorized"}),401
    db    = get_db()
    alert = db.execute("SELECT * FROM alerts WHERE id=?", (aid,)).fetchone()
    if not alert: return jsonify({"error":"Alert not found."}),404
    if session.get("role")=="admin" or alert["user_id"]==uid():
        db.execute("DELETE FROM alerts WHERE id=?", (aid,))
        db.commit()
        return jsonify({"message":"Alert deleted."})
    return jsonify({"error":"You can only delete your own alerts."}),403

@app.route("/api/alerts/<int:aid>/approve", methods=["POST"])
def approve_alert(aid):
    if not uid() or session.get("role")!="admin": return jsonify({"error":"Unauthorized"}),401
    db = get_db()
    db.execute("UPDATE alerts SET status='approved' WHERE id=?", (aid,))
    db.commit()
    return jsonify({"message":"Alert approved and now live."})

@app.route("/api/alerts/<int:aid>/decline", methods=["POST"])
def decline_alert(aid):
    if not uid() or session.get("role")!="admin": return jsonify({"error":"Unauthorized"}),401
    db = get_db()
    db.execute("UPDATE alerts SET status='declined' WHERE id=?", (aid,))
    db.commit()
    return jsonify({"message":"Alert declined."})

# ── Reports ───────────────────────────────────────────────────────────────────
@app.route("/api/reports", methods=["GET"])
def get_reports():
    if not uid(): return jsonify({"error":"Unauthorized"}),401
    db   = get_db()
    rows = db.execute(
        "SELECT r.id,r.location,r.detail,r.created_at,r.user_id,r.room_id,u.name as student_name "
        "FROM reports r LEFT JOIN users u ON r.user_id=u.id ORDER BY r.created_at DESC"
    ).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/reports", methods=["POST"])
def post_report():
    if not uid(): return jsonify({"error":"Unauthorized"}),401
    data     = request.get_json()
    location = data.get("location","").strip()
    detail   = data.get("detail","").strip()
    room_id  = data.get("room_id","").strip()
    floor    = data.get("floor","").strip()
    pos_x    = data.get("pos_x", None)
    pos_y    = data.get("pos_y", None)
    if not location: return jsonify({"error":"Location is required."}),400
    db = get_db()
    db.execute(
        "INSERT INTO reports (user_id,location,detail,created_at,room_id,floor,pos_x,pos_y) VALUES (?,?,?,?,?,?,?,?)",
        (uid(), location, detail, now(), room_id, floor, pos_x, pos_y)
    )
    db.commit()
    confirmed = False
    if room_id:
        count = db.execute(
            "SELECT COUNT(DISTINCT user_id) FROM reports WHERE room_id=?", (room_id,)
        ).fetchone()[0]
        if count >= 3:
            confirmed = True
    return jsonify({"message":"Hazard report submitted.", "room_id":room_id, "confirmed":confirmed})

@app.route("/api/reports/confirmed", methods=["GET"])
def get_confirmed_reports():
    db   = get_db()
    rows = db.execute(
        "SELECT DISTINCT room_id, floor, detail, pos_x, pos_y FROM reports "
        "WHERE room_id != '' GROUP BY room_id HAVING COUNT(DISTINCT user_id) >= 3"
    ).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/reports/my", methods=["GET"])
def get_my_reports():
    if not uid(): return jsonify({"error":"Unauthorized"}),401
    db   = get_db()
    rows = db.execute(
        "SELECT id,room_id,location,detail,created_at FROM reports WHERE user_id=? ORDER BY created_at DESC",
        (uid(),)
    ).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/reports/<int:rid>/mine", methods=["DELETE"])
def delete_my_report(rid):
    if not uid(): return jsonify({"error":"Unauthorized"}),401
    db     = get_db()
    report = db.execute("SELECT * FROM reports WHERE id=?", (rid,)).fetchone()
    if not report: return jsonify({"error":"Not found."}),404
    if report["user_id"]!=uid():
        return jsonify({"error":"You can only delete your own reports."}),403
    db.execute("DELETE FROM reports WHERE id=?", (rid,))
    db.commit()
    return jsonify({"message":"Report removed."})

@app.route("/api/reports/<int:rid>", methods=["DELETE"])
def delete_report(rid):
    if not uid() or session.get("role")!="admin": return jsonify({"error":"Unauthorized"}),401
    db = get_db()
    db.execute("DELETE FROM reports WHERE id=?", (rid,))
    db.commit()
    return jsonify({"message":"Report deleted."})

# ── Safety ────────────────────────────────────────────────────────────────────
@app.route("/api/safe", methods=["POST"])
def mark_safe():
    if not uid(): return jsonify({"error":"Unauthorized"}),401
    db = get_db()
    db.execute("INSERT OR REPLACE INTO safety_status (user_id,is_safe,updated_at) VALUES (?,1,?)", (uid(),now()))
    db.commit()
    return jsonify({"message":"You have been marked as safe."})

@app.route("/api/safe/admin", methods=["POST"])
def admin_mark_safe():
    if not uid() or session.get("role")!="admin": return jsonify({"error":"Unauthorized"}),401
    user_id = request.get_json().get("user_id")
    db = get_db()
    db.execute("INSERT OR REPLACE INTO safety_status (user_id,is_safe,updated_at) VALUES (?,1,?)", (user_id,now()))
    db.commit()
    return jsonify({"message":"Student marked as safe."})

@app.route("/api/safe/reset", methods=["POST"])
def reset_safe():
    if not uid() or session.get("role")!="admin": return jsonify({"error":"Unauthorized"}),401
    db = get_db()
    db.execute("UPDATE safety_status SET is_safe=0, updated_at=?", (now(),))
    db.commit()
    return jsonify({"message":"All safety statuses reset."})

@app.route("/api/safety", methods=["GET"])
def get_safety():
    if not uid() or session.get("role")!="admin": return jsonify({"error":"Unauthorized"}),401
    db   = get_db()
    rows = db.execute(
        "SELECT u.id,u.name,u.student_id,COALESCE(s.is_safe,0) as is_safe,s.updated_at "
        "FROM users u LEFT JOIN safety_status s ON u.id=s.user_id "
        "WHERE u.role IN ('student','staff','teacher')"
    ).fetchall()
    return jsonify([dict(r) for r in rows])

# ── SOS ───────────────────────────────────────────────────────────────────────
@app.route("/api/sos", methods=["POST"])
def send_sos():
    if not uid(): return jsonify({"error":"Unauthorized"}),401
    db = get_db()
    existing = db.execute(
        "SELECT id FROM sos_alerts WHERE user_id=? AND resolved=0 AND cancelled=0",(uid(),)
    ).fetchone()
    if existing: return jsonify({"error":"You already have an active SOS."}),409
    db.execute(
        "INSERT INTO sos_alerts (user_id,sender,created_at,resolved,responded,cancelled) VALUES (?,?,?,0,0,0)",
        (uid(), session.get("name","Unknown"), now())
    )
    db.commit()
    return jsonify({"message":"SOS sent."})

@app.route("/api/sos", methods=["GET"])
def get_sos():
    if not uid() or session.get("role")!="admin": return jsonify({"error":"Unauthorized"}),401
    db   = get_db()
    rows = db.execute(
        "SELECT s.*,u.name as student_name,u.email,u.student_id FROM sos_alerts s "
        "JOIN users u ON s.user_id=u.id WHERE s.cancelled=0 ORDER BY s.created_at DESC"
    ).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/sos/my", methods=["GET"])
def my_sos():
    if not uid(): return jsonify({"error":"Unauthorized"}),401
    db  = get_db()
    row = db.execute(
        "SELECT * FROM sos_alerts WHERE user_id=? AND cancelled=0 ORDER BY created_at DESC LIMIT 1",(uid(),)
    ).fetchone()
    if not row: return jsonify({"active":False})
    return jsonify(dict(row)|{"active":True})

@app.route("/api/sos/<int:sid>/respond", methods=["POST"])
def respond_sos(sid):
    if not uid() or session.get("role")!="admin": return jsonify({"error":"Unauthorized"}),401
    db  = get_db()
    sos = db.execute("SELECT * FROM sos_alerts WHERE id=?", (sid,)).fetchone()
    if not sos: return jsonify({"error":"Not found."}),404
    if sos["responded"]: return jsonify({"error":"Already responded."}),400
    db.execute(
        "UPDATE sos_alerts SET responded=1,responded_at=?,notification=? WHERE id=?",
        (now(),"A safety officer has been deployed and is on the way to your location. Please stay calm.",sid)
    )
    db.commit()
    return jsonify({"message":"Student notified."})

@app.route("/api/sos/<int:sid>/resolve", methods=["POST"])
def resolve_sos(sid):
    if not uid() or session.get("role")!="admin": return jsonify({"error":"Unauthorized"}),401
    db  = get_db()
    sos = db.execute("SELECT * FROM sos_alerts WHERE id=?", (sid,)).fetchone()
    if not sos: return jsonify({"error":"Not found."}),404
    if not sos["responded"]: return jsonify({"error":"Must mark responded before marking safe."}),400
    db.execute(
        "UPDATE sos_alerts SET resolved=1,notification=? WHERE id=?",
        ("You have been marked as safe by a safety officer. Please proceed to the assembly area.",sid)
    )
    db.execute(
        "INSERT OR REPLACE INTO safety_status (user_id,is_safe,updated_at) VALUES (?,1,?)",
        (sos["user_id"],now())
    )
    db.commit()
    return jsonify({"message":"Student marked as safe."})

@app.route("/api/sos/<int:sid>/cancel", methods=["POST"])
def cancel_sos(sid):
    if not uid(): return jsonify({"error":"Unauthorized"}),401
    db  = get_db()
    sos = db.execute("SELECT * FROM sos_alerts WHERE id=?", (sid,)).fetchone()
    if not sos: return jsonify({"error":"Not found."}),404
    if sos["user_id"]!=uid(): return jsonify({"error":"Not your SOS."}),403
    if sos["responded"]: return jsonify({"error":"Cannot cancel — officer already on the way."}),400
    db.execute("UPDATE sos_alerts SET cancelled=1 WHERE id=?", (sid,))
    db.commit()
    return jsonify({"message":"SOS cancelled."})

# ── Room Status (Admin Override) ──────────────────────────────────────────────
@app.route("/api/room-status", methods=["GET"])
def get_room_status():
    db   = get_db()
    rows = db.execute("SELECT * FROM room_status").fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/room-status", methods=["POST"])
def set_room_status():
    if not uid() or session.get("role")!="admin":
        return jsonify({"error":"Unauthorized"}),401
    data    = request.get_json()
    room_id = data.get("room_id","").strip()
    floor   = data.get("floor","").strip()
    status  = data.get("status","passable")
    if not room_id: return jsonify({"error":"room_id required"}),400
    db = get_db()
    db.execute(
        "INSERT INTO room_status (room_id,floor,status,set_by,updated_at) VALUES (?,?,?,?,?) "
        "ON CONFLICT(room_id) DO UPDATE SET status=?,set_by=?,updated_at=?",
        (room_id, floor, status, uid(), now(), status, uid(), now())
    )
    db.commit()
    return jsonify({"message":"Room status updated.","room_id":room_id,"status":status})

@app.route("/api/room-status/reset", methods=["POST"])
def reset_room_status():
    if not uid() or session.get("role")!="admin":
        return jsonify({"error":"Unauthorized"}),401
    db = get_db()
    db.execute("UPDATE room_status SET status='passable'")
    db.commit()
    return jsonify({"message":"All rooms reset to passable."})

# ── Evacuation Status ─────────────────────────────────────────────────────────
@app.route("/api/evacuating", methods=["GET"])
def get_evacuating():
    if not uid() or session.get("role")!="admin":
        return jsonify({"error":"Unauthorized"}),401
    db   = get_db()
    rows = db.execute(
        "SELECT u.id,u.name,u.student_id,s.room,s.floor,s.updated_at "
        "FROM evacuation_status s JOIN users u ON s.user_id=u.id "
        "WHERE s.is_evacuating=1"
    ).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/evacuating", methods=["POST"])
def set_evacuating():
    if not uid(): return jsonify({"error":"Unauthorized"}),401
    data  = request.get_json()
    room  = data.get("room","")
    floor = data.get("floor","")
    db    = get_db()
    db.execute(
        "INSERT INTO evacuation_status (user_id,room,floor,is_evacuating,updated_at) VALUES (?,?,?,1,?) "
        "ON CONFLICT(user_id) DO UPDATE SET room=?,floor=?,is_evacuating=1,updated_at=?",
        (uid(), room, floor, now(), room, floor, now())
    )
    db.commit()
    return jsonify({"message":"Evacuation status updated."})

@app.route("/api/evacuating/stop", methods=["POST"])
def stop_evacuating():
    if not uid(): return jsonify({"error":"Unauthorized"}),401
    db = get_db()
    db.execute(
        "INSERT INTO evacuation_status (user_id,room,floor,is_evacuating,updated_at) VALUES (?,?,?,0,?) "
        "ON CONFLICT(user_id) DO UPDATE SET is_evacuating=0,updated_at=?",
        (uid(),"","",now(),now())
    )
    db.commit()
    return jsonify({"message":"Evacuation stopped."})

# ── Schedule ──────────────────────────────────────────────────────────────────
@app.route("/api/schedule/upload", methods=["POST"])
def upload_schedule():
    if not uid() or session.get("role") != "admin":
        return jsonify({"error": "Unauthorized"}), 401
    data = request.get_json()
    b64  = data.get("file", "")
    if not b64:
        return jsonify({"error": "No file provided."}), 400
    raw  = base64.b64decode(b64)
    wb   = openpyxl.load_workbook(filename=io.BytesIO(raw))
    ws   = wb.active
    db   = get_db()
    added, skipped = 0, 0
    headers = [str(c.value).strip().lower() if c.value else '' for c in ws[1]]
    try:
        idx_sid   = headers.index('student_id')
        idx_subj  = headers.index('subject')
        idx_room  = headers.index('room')
        idx_day   = headers.index('day')
        idx_start = headers.index('time_start')
        idx_end   = headers.index('time_end')
    except ValueError:
        return jsonify({"error": "Excel must have columns: student_id, subject, room, day, time_start, time_end"}), 400
    for row in ws.iter_rows(min_row=2, values_only=True):
        try:
            sid     = str(row[idx_sid]).strip()   if row[idx_sid]   else ''
            subject = str(row[idx_subj]).strip()  if row[idx_subj]  else ''
            room    = str(row[idx_room]).strip()  if row[idx_room]  else ''
            day     = str(row[idx_day]).strip()   if row[idx_day]   else ''
            t_start = str(row[idx_start]).strip() if row[idx_start] else ''
            t_end   = str(row[idx_end]).strip()   if row[idx_end]   else ''
            if not all([sid, subject, room, day, t_start, t_end]):
                skipped += 1; continue
            if ':' not in t_start: t_start = t_start + ':00'
            if ':' not in t_end:   t_end   = t_end   + ':00'
            db.execute(
                "INSERT INTO schedules (student_id,subject,room,day,time_start,time_end) VALUES (?,?,?,?,?,?)",
                (sid, subject, room, day.upper(), t_start, t_end)
            )
            added += 1
        except Exception:
            skipped += 1
    db.commit()
    return jsonify({"message": f"Schedule uploaded. Added {added} entries, skipped {skipped}."})

@app.route("/api/schedule/current", methods=["GET"])
def get_current_schedule():
    if not uid(): return jsonify({"found": False, "message": "Unauthorized"}), 200
    from datetime import datetime
    db   = get_db()
    user = db.execute("SELECT student_id FROM users WHERE id=?", (uid(),)).fetchone()
    if not user: return jsonify({"found": False, "message": "User not found"}), 200
    sid  = user["student_id"]
    if not sid: return jsonify({"found": False, "message": "No student ID"}), 200
    now_dt   = datetime.now()
    now_time = now_dt.strftime("%H:%M")
    day_map  = {0:'M', 1:'T', 2:'W', 3:'TH', 4:'F', 5:'S', 6:'S'}
    today    = day_map[now_dt.weekday()]
    rows     = db.execute("SELECT * FROM schedules WHERE student_id=?", (sid,)).fetchall()
    for row in rows:
        if day_matches(today, row["day"].upper()):
            if row["time_start"] <= now_time <= row["time_end"]:
                return jsonify({
                    "found": True,
                    "student_id": sid,
                    "subject": row["subject"],
                    "room": row["room"],
                    "floor": get_floor_from_room(row["room"]),
                    "time_start": row["time_start"],
                    "time_end": row["time_end"],
                    "day": row["day"]
                })
    return jsonify({"found": False, "message": "No class scheduled at this time."})

def day_matches(today, day_code):
    day_code = day_code.upper().strip()
    patterns = {
        'MWF':  ['M','W','F'],
        'TTH':  ['T','TH'],
        'MW':   ['M','W'],
        'TF':   ['T','F'],
        'WF':   ['W','F'],
        'MF':   ['M','F'],
        'MTH':  ['M','TH'],
        'WTH':  ['W','TH'],
        'TW':   ['T','W'],
        'S':    ['S'],
        'M':    ['M'],
        'T':    ['T'],
        'W':    ['W'],
        'TH':   ['TH'],
        'F':    ['F'],
    }
    if day_code in patterns:
        return today in patterns[day_code]
    return today == day_code

def get_floor_from_room(room):
    import re
    room = room.upper().strip()
    m = re.match(r'BE(\d)\d{2}', room)
    if m:
        f = m.group(1)
        if f == '1': return 'Ground'
        if f == '2': return 'Second'
        if f == '3': return 'Third'
    ground = ['CANTEEN','PUMPROOM','EXIT']
    second = ['SCIENCELAB','GUIDANCEOFFICE','UTILITYROOM']
    third  = ['FUNCTIONHALL','LIBRARY','INTERNETSECTION','FACULTY','STUDENTLOUNGE']
    r = room.replace(' ','').replace('_','')
    for x in ground:
        if x in r: return 'Ground'
    for x in second:
        if x in r: return 'Second'
    for x in third:
        if x in r: return 'Third'
    return 'Ground'

@app.route("/api/schedule/all", methods=["GET"])
def get_all_schedules():
    if not uid() or session.get("role") != "admin":
        return jsonify({"error": "Unauthorized"}), 401
    db   = get_db()
    rows = db.execute(
        "SELECT s.*, u.name as student_name FROM schedules s "
        "LEFT JOIN users u ON s.student_id = u.student_id "
        "ORDER BY s.student_id, s.day, s.time_start"
    ).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route("/api/schedule/clear", methods=["POST"])
def clear_schedules():
    if not uid() or session.get("role") != "admin":
        return jsonify({"error": "Unauthorized"}), 401
    db = get_db()
    db.execute("DELETE FROM schedules")
    db.commit()
    return jsonify({"message": "All schedules cleared."})

@app.route("/api/schedule/<int:sid>", methods=["PATCH"])
def update_schedule(sid):
    if not uid() or session.get("role") != "admin":
        return jsonify({"error": "Unauthorized"}), 401
    data    = request.get_json()
    db      = get_db()
    row     = db.execute("SELECT * FROM schedules WHERE id=?", (sid,)).fetchone()
    if not row: return jsonify({"error": "Not found."}), 404
    subject    = data.get("subject",    row["subject"])
    room       = data.get("room",       row["room"])
    day        = data.get("day",        row["day"])
    time_start = data.get("time_start", row["time_start"])
    time_end   = data.get("time_end",   row["time_end"])
    db.execute(
        "UPDATE schedules SET subject=?,room=?,day=?,time_start=?,time_end=? WHERE id=?",
        (subject, room, day.upper(), time_start, time_end, sid)
    )
    db.commit()
    return jsonify({"message": "Schedule updated."})

@app.route("/api/schedule/<int:sid>", methods=["DELETE"])
def delete_schedule(sid):
    if not uid() or session.get("role") != "admin":
        return jsonify({"error": "Unauthorized"}), 401
    db = get_db()
    db.execute("DELETE FROM schedules WHERE id=?", (sid,))
    db.commit()
    return jsonify({"message": "Schedule entry deleted."})

@app.route("/api/schedule", methods=["POST"])
def add_schedule():
    if not uid() or session.get("role") != "admin":
        return jsonify({"error": "Unauthorized"}), 401
    data       = request.get_json()
    student_id = data.get("student_id","").strip()
    subject    = data.get("subject","").strip()
    room       = data.get("room","").strip()
    day        = data.get("day","").strip().upper()
    time_start = data.get("time_start","").strip()
    time_end   = data.get("time_end","").strip()
    if not all([student_id, subject, room, day, time_start, time_end]):
        return jsonify({"error": "All fields are required."}), 400
    db = get_db()
    db.execute(
        "INSERT INTO schedules (student_id,subject,room,day,time_start,time_end) VALUES (?,?,?,?,?,?)",
        (student_id, subject, room, day, time_start, time_end)
    )
    db.commit()
    return jsonify({"message": "Schedule entry added."})

if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)